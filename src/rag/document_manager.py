"""
文档管理模块
处理文档上传、解析、分块和存储
"""

import asyncio
from pathlib import Path
from typing import List, Dict, Any, Optional, Union
import hashlib
import json
from datetime import datetime

# 文档处理库
import PyPDF2
from docx import Document
import pandas as pd
from unstructured.partition.auto import partition

# 向量化和存储
from sentence_transformers import SentenceTransformer
import faiss
import numpy as np

from src.config.settings import Settings
from src.utils.logger import setup_logger

class DocumentManager:
    """文档管理器"""
    
    def __init__(self, settings: Settings):
        self.settings = settings
        self.logger = setup_logger(settings)
        
        # 初始化嵌入模型（添加网络超时和错误处理）
        self.embedding_model = self._initialize_embedding_model()
        
        # 初始化向量数据库
        self.vector_db = None
        self.document_metadata = {}
        self.chunks_data = []
        
        # 加载现有数据
        self._load_existing_data()
    
    def _get_local_model_path(self, model_name: str) -> Optional[str]:
        """获取本地模型路径"""
        # 项目本地模型目录
        project_models_dir = self.settings.project_root / "models" / "embeddings"
        
        # 检查项目本地模型
        if project_models_dir.exists():
            # 检查 HuggingFace 缓存格式的模型
            model_cache_name = f"models--sentence-transformers--{model_name.replace('/', '--')}"
            model_cache_path = project_models_dir / model_cache_name
            
            if model_cache_path.exists():
                self.logger.info(f"找到项目本地模型: {model_cache_path}")
                return str(project_models_dir)
        
        # 检查系统默认缓存
        try:
            import sentence_transformers
            default_cache = sentence_transformers.util.get_cache_folder()
            cache_path = Path(default_cache)
            
            model_cache_name = f"models--sentence-transformers--{model_name.replace('/', '--')}"
            model_cache_path = cache_path / model_cache_name
            
            if model_cache_path.exists():
                self.logger.info(f"找到系统缓存模型: {model_cache_path}")
                return str(cache_path)
        except Exception as e:
            self.logger.debug(f"检查系统缓存失败: {e}")
        
        return None
    
    def _initialize_embedding_model(self):
        """初始化嵌入模型，优先使用本地模型"""
        try:
            self.logger.info(f"正在初始化嵌入模型: {self.settings.embedding_model}")
            
            # 1. 首先检查本地模型路径
            local_cache_path = self._get_local_model_path(self.settings.embedding_model)
            if local_cache_path:
                try:
                    # 设置环境变量指向本地缓存
                    import os
                    original_cache = os.environ.get('SENTENCE_TRANSFORMERS_HOME')
                    os.environ['SENTENCE_TRANSFORMERS_HOME'] = local_cache_path
                    
                    model = SentenceTransformer(self.settings.embedding_model, local_files_only=True)
                    
                    # 恢复原始环境变量
                    if original_cache:
                        os.environ['SENTENCE_TRANSFORMERS_HOME'] = original_cache
                    elif 'SENTENCE_TRANSFORMERS_HOME' in os.environ:
                        del os.environ['SENTENCE_TRANSFORMERS_HOME']
                    
                    self.logger.info("嵌入模型初始化成功（本地模型）")
                    return model
                except Exception as local_error:
                    self.logger.warning(f"本地模型加载失败: {local_error}")
            
            # 2. 尝试默认离线模式
            try:
                model = SentenceTransformer(self.settings.embedding_model, local_files_only=True)
                self.logger.info("嵌入模型初始化成功（离线模式）")
                return model
            except Exception as offline_error:
                self.logger.info(f"离线模式失败: {offline_error}")
            
            # 3. 尝试备用本地模型
            backup_models = ["all-MiniLM-L6-v2", "paraphrase-multilingual-MiniLM-L12-v2"]
            for backup_model in backup_models:
                if backup_model == self.settings.embedding_model:
                    continue
                    
                backup_cache_path = self._get_local_model_path(backup_model)
                if backup_cache_path:
                    try:
                        import os
                        original_cache = os.environ.get('SENTENCE_TRANSFORMERS_HOME')
                        os.environ['SENTENCE_TRANSFORMERS_HOME'] = backup_cache_path
                        
                        model = SentenceTransformer(backup_model, local_files_only=True)
                        
                        # 恢复原始环境变量
                        if original_cache:
                            os.environ['SENTENCE_TRANSFORMERS_HOME'] = original_cache
                        elif 'SENTENCE_TRANSFORMERS_HOME' in os.environ:
                            del os.environ['SENTENCE_TRANSFORMERS_HOME']
                        
                        self.logger.info(f"备用本地模型初始化成功: {backup_model}")
                        return model
                    except Exception as backup_error:
                        self.logger.debug(f"备用模型 {backup_model} 加载失败: {backup_error}")
            
            # 4. 最后尝试在线下载（带超时）
            try:
                self.logger.info("尝试在线下载模型...")
                import socket
                original_timeout = socket.getdefaulttimeout()
                socket.setdefaulttimeout(10)  # 10秒超时
                
                model = SentenceTransformer(self.settings.embedding_model)
                
                # 恢复原始超时设置
                socket.setdefaulttimeout(original_timeout)
                
                self.logger.info("嵌入模型初始化成功（在线下载）")
                return model
            except Exception as online_error:
                self.logger.warning(f"在线下载失败: {online_error}")
                
                # 恢复原始超时设置
                if 'original_timeout' in locals():
                    socket.setdefaulttimeout(original_timeout)
            
        except Exception as e:
            self.logger.error(f"模型初始化过程出错: {e}")
        
        # 最后的备用方案：虚拟模型
        self.logger.warning("使用虚拟嵌入模型作为最后备用方案")
        return self._create_dummy_embedding_model()
    
    def _create_dummy_embedding_model(self):
        """创建一个虚拟的嵌入模型作为备用方案"""
        class DummyEmbeddingModel:
            def encode(self, texts, **kwargs):
                # 返回随机向量作为嵌入
                import numpy as np
                if isinstance(texts, str):
                    texts = [texts]
                return np.random.rand(len(texts), 384).astype(np.float32)
        
        return DummyEmbeddingModel()
    
    def _load_existing_data(self):
        """加载现有的向量数据库和元数据"""
        try:
            # 加载FAISS索引
            index_path = self.settings.vector_db_dir / "faiss_index.bin"
            if index_path.exists():
                self.vector_db = faiss.read_index(str(index_path))
                self.logger.info(f"加载现有FAISS索引，包含 {self.vector_db.ntotal} 个向量")
            
            # 加载元数据
            metadata_path = self.settings.vector_db_dir / "metadata.json"
            if metadata_path.exists():
                with open(metadata_path, 'r', encoding='utf-8') as f:
                    data = json.load(f)
                    self.document_metadata = data.get('documents', {})
                    self.chunks_data = data.get('chunks', [])
                    self.logger.info(f"加载 {len(self.document_metadata)} 个文档的元数据")
        except Exception as e:
            self.logger.error(f"加载现有数据失败: {e}")
    
    def _save_data(self):
        """保存向量数据库和元数据"""
        try:
            # 保存FAISS索引
            if self.vector_db:
                index_path = self.settings.vector_db_dir / "faiss_index.bin"
                faiss.write_index(self.vector_db, str(index_path))
            
            # 保存元数据
            metadata_path = self.settings.vector_db_dir / "metadata.json"
            with open(metadata_path, 'w', encoding='utf-8') as f:
                json.dump({
                    'documents': self.document_metadata,
                    'chunks': self.chunks_data
                }, f, ensure_ascii=False, indent=2)
            
            self.logger.info("数据保存成功")
        except Exception as e:
            self.logger.error(f"保存数据失败: {e}")
    
    def _extract_text_from_file(self, file_path: Path) -> str:
        """从文件中提取文本"""
        try:
            if file_path.suffix.lower() == '.pdf':
                return self._extract_from_pdf(file_path)
            elif file_path.suffix.lower() == '.docx':
                return self._extract_from_docx(file_path)
            elif file_path.suffix.lower() in ['.xlsx', '.xls']:
                return self._extract_from_excel(file_path)
            elif file_path.suffix.lower() == '.txt':
                with open(file_path, 'r', encoding='utf-8') as f:
                    return f.read()
            else:
                # 使用unstructured库处理其他格式
                elements = partition(str(file_path))
                return '\n'.join([str(element) for element in elements])
        except Exception as e:
            self.logger.error(f"提取文件 {file_path} 的文本失败: {e}")
            return ""
    
    def _extract_from_pdf(self, file_path: Path) -> str:
        """从PDF提取文本，增强版本"""
        try:
            # 首先尝试使用PyPDF2
            text = ""
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                for page in pdf_reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
            
            # 如果PyPDF2提取的文本太少，尝试使用unstructured库
            if len(text.strip()) < 100:
                self.logger.info(f"PyPDF2提取的文本较少，尝试使用unstructured库: {file_path}")
                try:
                    from unstructured.partition.pdf import partition_pdf
                    elements = partition_pdf(str(file_path))
                    unstructured_text = '\n'.join([str(element) for element in elements])
                    
                    # 如果unstructured提取的文本更多，使用它
                    if len(unstructured_text.strip()) > len(text.strip()):
                        text = unstructured_text
                except Exception as e:
                    self.logger.warning(f"使用unstructured提取PDF失败: {e}")
            
            return text
        except Exception as e:
            self.logger.error(f"PDF文本提取失败: {e}")
            # 最后尝试使用unstructured作为备选
            try:
                from unstructured.partition.pdf import partition_pdf
                elements = partition_pdf(str(file_path))
                return '\n'.join([str(element) for element in elements])
            except Exception as backup_error:
                self.logger.error(f"备选PDF提取也失败: {backup_error}")
                return ""
    
    def _extract_from_docx(self, file_path: Path) -> str:
        """从DOCX提取文本，增强版本，包括表格和格式化内容"""
        try:
            doc = Document(file_path)
            text = ""
            
            # 提取段落文本
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text += paragraph.text.strip() + "\n"
            
            # 提取表格内容
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        # 获取单元格中的所有段落
                        cell_text = '\n'.join([p.text.strip() for p in cell.paragraphs if p.text.strip()])
                        if cell_text:
                            row_text.append(cell_text)
                    if row_text:
                        text += ' | '.join(row_text) + "\n"
            
            # 如果提取的文本太少，尝试使用unstructured库
            if len(text.strip()) < 100:
                self.logger.info(f"python-docx提取的文本较少，尝试使用unstructured库: {file_path}")
                try:
                    from unstructured.partition.docx import partition_docx
                    elements = partition_docx(str(file_path))
                    unstructured_text = '\n'.join([str(element) for element in elements])
                    
                    # 如果unstructured提取的文本更多，使用它
                    if len(unstructured_text.strip()) > len(text.strip()):
                        text = unstructured_text
                except Exception as e:
                    self.logger.warning(f"使用unstructured提取DOCX失败: {e}")
            
            return text
        except Exception as e:
            self.logger.error(f"DOCX文本提取失败: {e}")
            # 最后尝试使用unstructured作为备选
            try:
                from unstructured.partition.docx import partition_docx
                elements = partition_docx(str(file_path))
                return '\n'.join([str(element) for element in elements])
            except Exception as backup_error:
                self.logger.error(f"备选DOCX提取也失败: {backup_error}")
                return ""
    
    def _extract_from_excel(self, file_path: Path) -> str:
        """从Excel提取文本"""
        df = pd.read_excel(file_path)
        return df.to_string()
    
    def _split_text_into_chunks(self, text: str, chunk_size: int = None, overlap: int = None) -> List[str]:
        """将文本分割成块"""
        chunk_size = chunk_size or self.settings.chunk_size
        overlap = overlap or self.settings.chunk_overlap
        
        if len(text) <= chunk_size:
            return [text]
        
        chunks = []
        start = 0
        
        while start < len(text):
            end = start + chunk_size
            chunk = text[start:end]
            
            # 尝试在句号或换行符处分割
            if end < len(text):
                last_period = chunk.rfind('。')
                last_newline = chunk.rfind('\n')
                split_point = max(last_period, last_newline)
                
                if split_point > start + chunk_size // 2:
                    chunk = text[start:start + split_point + 1]
                    end = start + split_point + 1
            
            chunks.append(chunk.strip())
            start = end - overlap
            
            if start >= len(text):
                break
        
        return [chunk for chunk in chunks if chunk.strip()]
    
    def _generate_document_id(self, content: str, title: str) -> str:
        """生成文档ID"""
        content_hash = hashlib.md5((content + title).encode()).hexdigest()
        return f"doc_{content_hash[:12]}"
    
    async def add_document(self, file_content: Union[bytes, str], filename: str, category: str = "其他") -> str:
        """添加文档到知识库"""
        try:
            # 保存文件
            file_path = self.settings.documents_dir / filename
            
            if isinstance(file_content, bytes):
                with open(file_path, 'wb') as f:
                    f.write(file_content)
            else:
                with open(file_path, 'w', encoding='utf-8') as f:
                    f.write(file_content)
            
            # 提取文本
            text = self._extract_text_from_file(file_path)
            if not text.strip():
                raise ValueError("无法从文件中提取文本")
            
            # 生成文档ID
            doc_id = self._generate_document_id(text, filename)
            
            # 检查是否已存在
            if doc_id in self.document_metadata:
                self.logger.info(f"文档 {filename} 已存在")
                return doc_id
            
            # 分割文本
            chunks = self._split_text_into_chunks(text)
            
            # 生成嵌入向量
            embeddings = self.embedding_model.encode(chunks)
            
            # 初始化或更新向量数据库
            if self.vector_db is None:
                dimension = embeddings.shape[1]
                self.vector_db = faiss.IndexFlatIP(dimension)  # 内积相似度
            
            # 添加向量到数据库
            start_idx = self.vector_db.ntotal
            self.vector_db.add(embeddings.astype('float32'))
            
            # 保存文档元数据
            self.document_metadata[doc_id] = {
                'filename': filename,
                'category': category,
                'upload_time': datetime.now().isoformat(),
                'chunk_count': len(chunks),
                'start_idx': start_idx,
                'end_idx': start_idx + len(chunks)
            }
            
            # 保存块数据
            for i, chunk in enumerate(chunks):
                self.chunks_data.append({
                    'doc_id': doc_id,
                    'chunk_idx': i,
                    'content': chunk,
                    'vector_idx': start_idx + i
                })
            
            # 保存数据
            self._save_data()
            
            self.logger.info(f"成功添加文档 {filename}，包含 {len(chunks)} 个文本块")
            return doc_id
            
        except Exception as e:
            self.logger.error(f"添加文档失败: {e}")
            raise
    

    
    async def search_documents(self, query: str, category: Optional[str] = None, top_k: int = None) -> List[Dict[str, Any]]:
        """搜索文档"""
        if not self.vector_db or self.vector_db.ntotal == 0:
            return []
        
        try:
            top_k = top_k or self.settings.top_k_results
            
            # 生成查询向量
            query_embedding = self.embedding_model.encode([query])
            
            # 搜索相似向量
            scores, indices = self.vector_db.search(query_embedding.astype('float32'), top_k)
            
            results = []
            for score, idx in zip(scores[0], indices[0]):
                if idx == -1:  # FAISS返回-1表示无效索引
                    continue
                
                # 找到对应的文本块
                chunk_info = None
                for chunk in self.chunks_data:
                    if chunk['vector_idx'] == idx:
                        chunk_info = chunk
                        break
                
                if chunk_info:
                    doc_info = self.document_metadata.get(chunk_info['doc_id'], {})
                    
                    # 类别过滤
                    if category and doc_info.get('category') != category:
                        continue
                    
                    # 相似度过滤
                    if score < self.settings.similarity_threshold:
                        continue
                    
                    results.append({
                        'content': chunk_info['content'],
                        'score': float(score),
                        'source': doc_info.get('filename', '未知'),
                        'category': doc_info.get('category', '其他'),
                        'doc_id': chunk_info['doc_id'],
                        'chunk_idx': chunk_info['chunk_idx']
                    })
            
            # 按相似度排序
            results.sort(key=lambda x: x['score'], reverse=True)
            return results
            
        except Exception as e:
            self.logger.error(f"搜索文档失败: {e}")
            return []
    
    async def get_documents_summary(self, document_type: Optional[str] = None) -> str:
        """获取文档摘要"""
        if not self.document_metadata:
            return "暂无上传的文档"
        
        summary = "文档库摘要:\n\n"
        
        # 按类别统计
        category_stats = {}
        for doc_id, doc_info in self.document_metadata.items():
            category = doc_info.get('category', '其他')
            if document_type and category != document_type:
                continue
            
            if category not in category_stats:
                category_stats[category] = []
            category_stats[category].append(doc_info)
        
        for category, docs in category_stats.items():
            summary += f"**{category}类文档** ({len(docs)}个):\n"
            for doc in docs[:5]:  # 只显示前5个
                summary += f"  - {doc['filename']} ({doc['chunk_count']}个文本块)\n"
            if len(docs) > 5:
                summary += f"  - ... 还有{len(docs)-5}个文档\n"
            summary += "\n"
        
        summary += f"总计: {len(self.document_metadata)}个文档，{len(self.chunks_data)}个文本块"
        return summary